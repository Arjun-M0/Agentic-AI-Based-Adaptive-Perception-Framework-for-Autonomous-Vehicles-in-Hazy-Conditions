from pathlib import Path
from typing import Dict, Iterable

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def _state_value(state, field: str):
    return getattr(state, field, None) if state is not None else None


def write_excel_report(run: Dict, output_file: str) -> str:
    """Write measured pipeline results and explicit unavailable fields to Excel."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Field", "Value"])
    summary_rows = [
        ("Input source", run["source"]),
        ("Processed images", len(run["results"])),
        ("Skipped images", len(run["skipped"])),
        (
            "Cold-start explorations",
            sum(1 for result in run["results"] if result.get("cold_start_exploration")),
        ),
        (
            "Knowledge-assisted results",
            sum(
                1
                for result in run["results"]
                if result.get("knowledge_assisted_latency_ms") is not None
            ),
        ),
        ("mAP", "Unavailable: no ground-truth annotations"),
        ("Precision", "Unavailable: no ground-truth annotations"),
        ("Recall", "Unavailable: no ground-truth annotations"),
    ]
    for row in summary_rows:
        summary.append(row)

    images = workbook.create_sheet("Per Image")
    images.append(
        [
            "Frame ID",
            "Filename",
            "Width",
            "Height",
            "Strategy",
            "Strategy Explanation",
            "Cold-Start Exploration",
            "Candidate Strategies Evaluated",
            "Candidate Latencies (ms)",
            "Cold-Start Exploration Total (ms)",
            "Previous Exploration Total (ms)",
            "Knowledge-Assisted Latency (ms)",
            "Latency Reduction (%)",
            "Haze Score",
            "Visibility Score",
            "Illumination Score",
            "Contrast Score",
            "CPU %",
            "RAM %",
            "GPU Available",
            "Restoration Latency (ms)",
            "Detection Latency (ms)",
            "Total Latency (ms)",
            "FPS",
            "Detection Count",
            "mAP",
            "Precision",
            "Recall",
        ]
    )
    detections = workbook.create_sheet("Detections")
    detections.append(["Frame ID", "Filename", "Class", "Confidence", "X1", "Y1", "X2", "Y2"])

    for result in run["results"]:
        environment = result["environment"]
        hardware = result["hardware"]
        decision = result["decision"]
        detection = result["detection"]
        evaluation = result["evaluation"]
        images.append(
            [
                result["frame_id"],
                result["filename"],
                result["width"],
                result["height"],
                decision.selected_strategy,
                decision.explanation,
                result.get("cold_start_exploration", False),
                result.get("candidate_count", 1),
                str(result.get("candidate_latencies_ms", {})),
                result.get("cold_start_exploration_total_latency_ms"),
                result.get("previous_exploration_latency_ms"),
                result.get("knowledge_assisted_latency_ms"),
                result.get("latency_reduction_percent"),
                _state_value(environment, "haze_score"),
                _state_value(environment, "visibility_score"),
                _state_value(environment, "illumination_score"),
                _state_value(environment, "contrast_score"),
                _state_value(hardware, "cpu_utilization"),
                _state_value(hardware, "ram_utilization"),
                _state_value(hardware, "gpu_available"),
                result["restoration_latency_ms"],
                detection["latency_ms"],
                evaluation["latency_ms"],
                evaluation["fps"],
                len(detection["detections"]),
                evaluation["mAP"],
                evaluation["precision"],
                evaluation["recall"],
            ]
        )
        for item in detection["detections"]:
            box = item["box"]
            detections.append(
                [
                    result["frame_id"],
                    result["filename"],
                    item["class_name"],
                    item["confidence"],
                    *box,
                ]
            )

    skipped = workbook.create_sheet("Skipped")
    skipped.append(["Filename", "Path", "Reason"])
    for item in run["skipped"]:
        skipped.append([item["filename"], item["path"], item["error"]])

    strategy_sheet = workbook.create_sheet("Strategy Summary")
    strategy_sheet.append(["Strategy", "Images", "Average Latency (ms)", "Average FPS"])
    strategies = {}
    for result in run["results"]:
        strategy = result["decision"].selected_strategy
        strategies.setdefault(strategy, []).append(result["evaluation"])
    for strategy, evaluations in sorted(strategies.items()):
        strategy_sheet.append(
            [
                strategy,
                len(evaluations),
                sum(item["latency_ms"] for item in evaluations) / len(evaluations),
                sum(item["fps"] for item in evaluations) / len(evaluations),
            ]
        )

    comparison = workbook.create_sheet("Knowledge Comparison")
    comparison.append(["Metric", "Before Knowledge", "After Knowledge"])
    result = run["results"][-1] if run["results"] else {}
    candidate_latencies = result.get("candidate_latencies_ms", {})
    previous_candidate_latencies = result.get("previous_candidate_latencies_ms")
    before_total = (
        result.get("previous_exploration_latency_ms")
        or result.get("cold_start_exploration_total_latency_ms")
    )
    after_latency = result.get("knowledge_assisted_latency_ms")
    comparison_rows = [
        ("Input", run["source"], run["source"]),
        (
            "Execution mode",
            "Cold-start exploration",
            "Knowledge-assisted execution",
        ),
        (
            "Strategies evaluated",
            3
            if result.get("cold_start_exploration")
            else len(previous_candidate_latencies or {}),
            result.get("candidate_count", 1),
        ),
        (
            "Strategies executed",
            "Bypass, DEA-Net, DehazeFormer"
            if result.get("cold_start_exploration")
            else ", ".join(previous_candidate_latencies or {})
            or "Stored first-run candidates",
            result.get("decision").selected_strategy if result.get("decision") else None,
        ),
        (
            "Candidate latency measurements (ms)",
            str(previous_candidate_latencies or candidate_latencies)
            if (previous_candidate_latencies or candidate_latencies)
            else "Stored in first-run report",
            str({result.get("decision").selected_strategy: after_latency})
            if result.get("decision") and after_latency is not None
            else "Unavailable",
        ),
        ("Total processing latency (ms)", before_total, after_latency),
        (
            "Latency reduction (%)",
            "Not applicable",
            result.get("latency_reduction_percent"),
        ),
    ]
    for row in comparison_rows:
        comparison.append(row)

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for column_cells in sheet.columns:
            width = min(
                60,
                max(len(str(cell.value or "")) for cell in column_cells) + 2,
            )
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

    destination = Path(output_file)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return str(destination)
